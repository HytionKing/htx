"""Generate the qualification-defense deck (native .pptx).

Slide 01 — Team management: business capability matrix.
Layout mirrors ppt/pages/01-team-capability-matrix.html (3:2 columns, equal heights).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# ---------- brand palette ----------
PRIMARY      = RGBColor(0xFF, 0x79, 0x00)
PRIMARY_SOFT = RGBColor(0xFF, 0xF3, 0xE8)
AUX          = RGBColor(0x55, 0xBA, 0xBE)
AUX_SOFT     = RGBColor(0xE6, 0xF5, 0xF5)
INK          = RGBColor(0x33, 0x33, 0x33)
INK_2        = RGBColor(0x60, 0x60, 0x60)
INK_3        = RGBColor(0xA8, 0xA8, 0xA8)
LINE         = RGBColor(0xEC, 0xEC, 0xEC)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
BG_SOFT      = RGBColor(0xFA, 0xFA, 0xFA)
BG_WARM      = RGBColor(0xFA, 0xF9, 0xF7)

FONT = "Microsoft YaHei"


# ---------- helpers ----------
def add_rect(slide, x, y, w, h, *, fill=None, line=None, line_w=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, corner=0.08, shadow=False):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = corner
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.text_frame.margin_left = shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = shp.text_frame.margin_bottom = 0
    shp.text_frame.word_wrap = True
    if not shadow:
        # remove default shadow
        sppr = shp._element.spPr
        from pptx.oxml.ns import qn
        for el in sppr.findall(qn("a:effectLst")):
            sppr.remove(el)
        eff = sppr.makeelement(qn("a:effectLst"), {})
        sppr.append(eff)
    return shp


def add_text(slide, x, y, w, h, text, *, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             spacing=1.15, letter_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if letter_spacing is not None:
            from pptx.oxml.ns import qn
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(letter_spacing))
    return tb


def add_rich(slide, x, y, w, h, runs, *, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: list of (text, dict(size, color, bold, font))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    for text, attrs in runs:
        r = p.add_run()
        r.text = text
        r.font.name = attrs.get("font", FONT)
        r.font.size = Pt(attrs.get("size", 14))
        r.font.bold = attrs.get("bold", False)
        if "color" in attrs:
            r.font.color.rgb = attrs["color"]
    return tb


def icon_box(slide, x, y, w, h, glyph, *, color=PRIMARY, bg=PRIMARY_SOFT, size=22):
    add_rect(slide, x, y, w, h, fill=bg, corner=0.22)
    add_text(slide, x, y, w, h, glyph, size=size, color=color, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ---------- build ----------
def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # Background — white default; soft corner gradients via faint rectangles
    add_rect(slide, -1.5, -1.5, 5.5, 5.5, fill=RGBColor(0xFF, 0xF3, 0xE8),
             shape=MSO_SHAPE.OVAL).fill.fore_color.rgb = RGBColor(0xFF, 0xF3, 0xE8)
    add_rect(slide, 11.5, -2.0, 5.5, 5.5, fill=RGBColor(0xFF, 0xF6, 0xEC),
             shape=MSO_SHAPE.OVAL)
    add_rect(slide, -1.5, 5.5, 5.5, 5.5, fill=RGBColor(0xEE, 0xF7, 0xF7),
             shape=MSO_SHAPE.OVAL)

    # ===== Header =====
    # Title bar
    add_rect(slide, 0.55, 0.45, 0.06, 0.7, fill=PRIMARY, corner=0.3)
    # Main title (mixed color)
    add_rich(slide, 0.75, 0.42, 9.0, 0.55,
             [("团队管理 · ", {"size": 30, "color": INK, "bold": True}),
              ("业务能力矩阵建设", {"size": 30, "color": PRIMARY, "bold": True})],
             anchor=MSO_ANCHOR.MIDDLE)
    # Subtitle
    add_text(slide, 0.75, 0.95, 9.0, 0.3,
             "TEAM CAPABILITY MATRIX   /   人岗适配 · 互备协同 · 动态调度",
             size=11, color=INK_2, letter_spacing=200, anchor=MSO_ANCHOR.MIDDLE)
    # Tag (right)
    tag = add_rect(slide, 10.55, 0.55, 1.85, 0.4, fill=WHITE,
                   line=PRIMARY, line_w=1.0, corner=0.5)
    add_text(slide, 10.55, 0.55, 1.85, 0.4, "PART 01 · 团队管理",
             size=10, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, letter_spacing=100)
    add_text(slide, 10.55, 1.0, 1.85, 0.25, "— 01 / 0X —",
             size=9, color=INK_3, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, letter_spacing=200)

    # Body coords
    body_top = 1.55
    body_bottom = 6.95
    body_h = body_bottom - body_top
    body_left = 0.55
    body_right = 12.78
    gap = 0.28
    avail_w = body_right - body_left - gap            # 11.95
    left_w = avail_w * 0.6                            # 7.17
    right_w = avail_w * 0.4                           # 4.78
    left_x = body_left
    right_x = body_left + left_w + gap

    row_gap = 0.22
    avail_h = body_h - row_gap                        # 5.18
    top_h = avail_h * (1.15 / 2.0)                    # 2.98
    bot_h = avail_h * (0.85 / 2.0)                    # 2.20
    top_y = body_top
    bot_y = body_top + top_h + row_gap

    # ===== LEFT TOP: Challenges =====
    pad = 0.26
    add_rect(slide, left_x, top_y, left_w, top_h, fill=WHITE, line=LINE, corner=0.05)
    # icon + title
    icon_box(slide, left_x + pad, top_y + pad, 0.42, 0.42, "!",
             color=PRIMARY, bg=PRIMARY_SOFT, size=20)
    add_text(slide, left_x + pad + 0.55, top_y + pad, 3.0, 0.42,
             "面临挑战", size=18, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, left_x + left_w - 2.0 - pad, top_y + pad, 2.0, 0.42,
             "CHALLENGES", size=9, color=INK_3, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE, letter_spacing=300)

    # bullets
    bullets = [
        [("业务复杂度持续提升，", {"size": 13, "color": INK}),
         ("多项目并行", {"size": 13, "color": PRIMARY, "bold": True}),
         ("开展，同一业务领域在多个项目内同步迭代", {"size": 13, "color": INK})],
        [("单人负责制", {"size": 13, "color": PRIMARY, "bold": True}),
         ("难以支撑业务增长，存在交付瓶颈与单点风险", {"size": 13, "color": INK})],
        [("人员请假 / 流动时，", {"size": 13, "color": INK}),
         ("知识无承接", {"size": 13, "color": PRIMARY, "bold": True}),
         ("，重点项目协同效率受损", {"size": 13, "color": INK})],
    ]
    by = top_y + pad + 0.65
    for runs in bullets:
        add_rect(slide, left_x + pad + 0.05, by + 0.16, 0.1, 0.1,
                 fill=PRIMARY, shape=MSO_SHAPE.OVAL)
        add_rich(slide, left_x + pad + 0.28, by, left_w - pad * 2 - 0.3, 0.5,
                 runs, spacing=1.4)
        by += 0.42

    # metrics row at bottom
    m_top = top_y + top_h - pad - 0.78
    m_w = (left_w - pad * 2 - 0.24) / 3
    metrics = [("5+", "条线", "并行业务"),
               ("3×", "", "迭代节奏"),
               ("100%", "", "人员互备")]
    for i, (num, unit, lbl) in enumerate(metrics):
        mx = left_x + pad + i * (m_w + 0.12)
        # dashed-look metric box (solid soft fill, thin line)
        add_rect(slide, mx, m_top, m_w, 0.78, fill=BG_WARM, line=LINE, corner=0.12)
        # number + unit
        add_rich(slide, mx, m_top + 0.1, m_w, 0.4,
                 [(num, {"size": 22, "color": PRIMARY, "bold": True})]
                 + ([(unit, {"size": 10, "color": INK_2, "bold": True})] if unit else []),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, mx, m_top + 0.5, m_w, 0.25, lbl,
                 size=10, color=INK_2, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, letter_spacing=150)

    # ===== LEFT BOTTOM: Approach =====
    add_rect(slide, left_x, bot_y, left_w, bot_h, fill=WHITE, line=LINE, corner=0.06)
    icon_box(slide, left_x + pad, bot_y + pad, 0.42, 0.42, "→",
             color=AUX, bg=AUX_SOFT, size=20)
    add_text(slide, left_x + pad + 0.55, bot_y + pad, 3.0, 0.42,
             "建设举措", size=18, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, left_x + left_w - 2.0 - pad, bot_y + pad, 2.0, 0.42,
             "APPROACH", size=9, color=INK_3, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE, letter_spacing=300)

    steps = [("STEP 1", "业务矩阵划分", "按业务特性切分领域，形成清晰的能力边界与协作单元。"),
             ("STEP 2", "人岗矩阵映射", "结合成员专长与意愿，主备搭配、能力对齐。"),
             ("STEP 3", "动态调度赋能", "重点业务允许多人参与，轮岗共建、沉淀知识。")]
    s_top = bot_y + pad + 0.65
    s_h = bot_h - pad - 0.7
    s_w = (left_w - pad * 2 - 0.24) / 3
    for i, (n, title, desc) in enumerate(steps):
        sx = left_x + pad + i * (s_w + 0.12)
        add_rect(slide, sx, s_top, s_w, s_h, fill=BG_SOFT, line=LINE, corner=0.1)
        # numbered badge
        add_rect(slide, sx + 0.2, s_top - 0.12, 0.7, 0.26, fill=PRIMARY, corner=0.4)
        add_text(slide, sx + 0.2, s_top - 0.12, 0.7, 0.26, n,
                 size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, letter_spacing=200)
        add_text(slide, sx + 0.2, s_top + 0.22, s_w - 0.4, 0.35,
                 title, size=14, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, sx + 0.2, s_top + 0.62, s_w - 0.4, s_h - 0.7,
                 desc, size=11, color=INK_2, spacing=1.55, anchor=MSO_ANCHOR.TOP)

    # ===== RIGHT TOP: Matrix =====
    rpad = 0.22
    add_rect(slide, right_x, top_y, right_w, top_h, fill=WHITE, line=LINE, corner=0.06)
    icon_box(slide, right_x + rpad, top_y + rpad, 0.42, 0.42, "▦",
             color=AUX, bg=AUX_SOFT, size=18)
    add_text(slide, right_x + rpad + 0.55, top_y + rpad, 2.2, 0.42,
             "业务 × 人员 矩阵", size=15, bold=True, color=INK,
             anchor=MSO_ANCHOR.MIDDLE)

    # legend (top right of card)
    lg_y = top_y + rpad + 0.06
    lg_x = right_x + right_w - rpad - 2.05
    # 主
    add_rect(slide, lg_x, lg_y, 0.22, 0.22, fill=PRIMARY, corner=0.25)
    add_text(slide, lg_x, lg_y, 0.22, 0.22, "主", size=8, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, lg_x + 0.26, lg_y, 0.5, 0.22, "主负责", size=9,
             color=INK_2, anchor=MSO_ANCHOR.MIDDLE)
    # 备
    lg_x2 = lg_x + 0.85
    add_rect(slide, lg_x2, lg_y, 0.22, 0.22, fill=WHITE, line=AUX, line_w=1.3, corner=0.25)
    add_text(slide, lg_x2, lg_y, 0.22, 0.22, "备", size=8, color=AUX,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, lg_x2 + 0.26, lg_y, 0.5, 0.22, "互备", size=9,
             color=INK_2, anchor=MSO_ANCHOR.MIDDLE)
    # —
    lg_x3 = lg_x2 + 0.75
    add_rect(slide, lg_x3, lg_y, 0.22, 0.22, fill=RGBColor(0xF3, 0xF3, 0xF3),
             line=LINE, corner=0.25)
    add_text(slide, lg_x3, lg_y, 0.22, 0.22, "—", size=8, color=INK_3,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, lg_x3 + 0.26, lg_y, 0.6, 0.22, "未覆盖", size=9,
             color=INK_2, anchor=MSO_ANCHOR.MIDDLE)

    # Matrix table
    mx0 = right_x + rpad
    my0 = top_y + rpad + 0.7
    mw = right_w - rpad * 2
    mh = top_h - rpad * 2 - 0.7
    # outer rounded border
    add_rect(slide, mx0, my0, mw, mh, fill=WHITE, line=LINE, corner=0.06)

    rows = [
        ("用户中心", "账户·权限·风控", ["主", "备", "—", "备", "—", "—"]),
        ("订单交易", "下单·履约·售后", ["备", "主", "主", "—", "备", "—"]),
        ("商品库存", "商品·库存·价格", ["—", "备", "主", "备", "—", "备"]),
        ("营销活动", "促销·投放",     ["—", "—", "备", "主", "主", "备"]),
        ("支付结算", "收银·对账",     ["备", "—", "—", "备", "备", "主"]),
    ]
    members = ["A", "B", "C", "D", "E", "F"]
    n_cols = 1 + len(members)
    n_rows = 1 + len(rows)
    # column widths
    label_col_w = 1.45
    cell_col_w = (mw - label_col_w) / len(members)
    header_h = 0.34
    body_h_t = mh - header_h
    row_h = body_h_t / len(rows)

    # header background
    add_rect(slide, mx0, my0, mw, header_h, fill=BG_SOFT, line=None, corner=0)
    add_text(slide, mx0 + 0.14, my0, label_col_w, header_h,
             "业务 / 成员", size=10, color=INK_2, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    for ci, m in enumerate(members):
        cx = mx0 + label_col_w + ci * cell_col_w
        add_text(slide, cx, my0, cell_col_w, header_h, m,
                 size=10, color=INK_2, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # body rows
    for ri, (name, sub, cells) in enumerate(rows):
        ry = my0 + header_h + ri * row_h
        # alt row label background (subtle)
        add_rect(slide, mx0, ry, label_col_w, row_h,
                 fill=RGBColor(0xFC, 0xFB, 0xFA), line=None, corner=0)
        add_rich(slide, mx0 + 0.14, ry + 0.04, label_col_w - 0.18, row_h - 0.08,
                 [(name + "\n", {"size": 11, "color": INK, "bold": True})],
                 anchor=MSO_ANCHOR.TOP)
        add_text(slide, mx0 + 0.14, ry + row_h * 0.55, label_col_w - 0.18,
                 row_h * 0.45, sub, size=8, color=INK_3, anchor=MSO_ANCHOR.TOP)

        for ci, mark in enumerate(cells):
            cx = mx0 + label_col_w + ci * cell_col_w
            # badge centered in cell
            bw = min(0.28, cell_col_w - 0.16)
            bh = min(0.28, row_h - 0.16)
            bx = cx + (cell_col_w - bw) / 2
            by = ry + (row_h - bh) / 2
            if mark == "主":
                add_rect(slide, bx, by, bw, bh, fill=PRIMARY, corner=0.3)
                add_text(slide, bx, by, bw, bh, "主", size=9, color=WHITE,
                         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            elif mark == "备":
                add_rect(slide, bx, by, bw, bh, fill=WHITE, line=AUX,
                         line_w=1.2, corner=0.3)
                add_text(slide, bx, by, bw, bh, "备", size=9, color=AUX,
                         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            else:
                add_rect(slide, bx, by, bw, bh,
                         fill=RGBColor(0xF3, 0xF3, 0xF3), corner=0.3)
                add_text(slide, bx, by, bw, bh, "—", size=9, color=INK_3,
                         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        # row divider
        if ri < len(rows) - 1:
            ln = slide.shapes.add_connector(1, Inches(mx0), Inches(ry + row_h),
                                            Inches(mx0 + mw), Inches(ry + row_h))
            ln.line.color.rgb = LINE
            ln.line.width = Pt(0.5)
    # header divider
    ln = slide.shapes.add_connector(1, Inches(mx0), Inches(my0 + header_h),
                                    Inches(mx0 + mw), Inches(my0 + header_h))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.5)

    # ===== RIGHT BOTTOM: Outcomes =====
    add_rect(slide, right_x, bot_y, right_w, bot_h, fill=WHITE, line=LINE, corner=0.06)
    icon_box(slide, right_x + rpad, bot_y + rpad, 0.42, 0.42, "↗",
             color=PRIMARY, bg=PRIMARY_SOFT, size=20)
    add_text(slide, right_x + rpad + 0.55, bot_y + rpad, 2.2, 0.42,
             "实施成效", size=15, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, right_x + right_w - 1.6 - rpad, bot_y + rpad, 1.6, 0.42,
             "OUTCOMES", size=9, color=INK_3, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE, letter_spacing=300)

    outs = [("交付效率", "并行吞吐", "↑ 40%", PRIMARY),
            ("单点风险", "关键业务", "0 阻塞", AUX),
            ("人才梯队", "互备覆盖", "100%", AUX),
            ("成长牵引", "人均覆盖", "+2 领域", PRIMARY)]
    ox_top = bot_y + rpad + 0.7
    ox_h = bot_h - rpad * 2 - 0.7
    oc_w = (right_w - rpad * 2 - 0.16) / 2
    oc_h = (ox_h - 0.14) / 2
    for i, (t, lbl, val, ac) in enumerate(outs):
        col = i % 2
        row = i // 2
        ox = right_x + rpad + col * (oc_w + 0.16)
        oy = ox_top + row * (oc_h + 0.14)
        add_rect(slide, ox, oy, oc_w, oc_h, fill=BG_SOFT, line=LINE, corner=0.12)
        # left color bar
        add_rect(slide, ox, oy, 0.07, oc_h, fill=ac, shape=MSO_SHAPE.RECTANGLE, corner=0)
        add_text(slide, ox + 0.22, oy + 0.1, oc_w - 0.3, 0.25,
                 t, size=10, color=INK_2, anchor=MSO_ANCHOR.MIDDLE, letter_spacing=150)
        add_rich(slide, ox + 0.22, oy + 0.4, oc_w - 0.3, oc_h - 0.45,
                 [(lbl + "  ", {"size": 13, "color": INK, "bold": True}),
                  (val, {"size": 15, "color": ac, "bold": True})],
                 anchor=MSO_ANCHOR.MIDDLE)

    # ===== Footer =====
    add_text(slide, 0.55, 7.12, 4.0, 0.25,
             "认证资格认证述职答辩 · 2026",
             size=9, color=INK_3, letter_spacing=200, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.5, 7.12, 4.3, 0.25,
             "TEAM MANAGEMENT  /  CAPABILITY MATRIX",
             size=9, color=INK_3, align=PP_ALIGN.RIGHT,
             letter_spacing=200, anchor=MSO_ANCHOR.MIDDLE)
    # thin divider
    ln = slide.shapes.add_connector(1, Inches(4.6), Inches(7.24),
                                    Inches(8.4), Inches(7.24))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.5)

    out_path = Path(__file__).parent / "dist" / "认证述职答辩.pptx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    build()
